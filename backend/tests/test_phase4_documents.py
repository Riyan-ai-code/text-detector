import io
import pytest
from fastapi.testclient import TestClient
import pypdf
import docx

from backend.app.main import app
from backend.app.services.document_forensics import document_forensics

client = TestClient(app)

def test_txt_zero_width_adversarial_detection():
    """Verify zero-width invisible character anomaly detection."""
    sample_text = (
        "This is an authentic essay\u200b with hidden\u200c prompt-injection tokens "
        "interspersed to fool classical lexical classifiers."
    )
    raw_bytes = sample_text.encode("utf-8")
    result = document_forensics.analyze_document_file("essay.txt", raw_bytes)
    assert len(result.formatting_anomalies) > 0
    font_anomalies = [a for a in result.formatting_anomalies if a.type == "font"]
    assert len(font_anomalies) > 0
    assert "Invisible adversarial unicode" in font_anomalies[0].description
    assert font_anomalies[0].severity == "HIGH"
    assert result.classification in ["POTENTIALLY_MANIPULATED", "AI_GENERATED"]

def test_citation_audit_missing_references():
    """Verify citation detection flags missing bibliography."""
    sample_text = (
        "Recent breakthroughs in deep neural modeling (Vaswani et al., 2017) and "
        "reinforcement learning [1] have achieved remarkable milestones."
    )
    raw_bytes = sample_text.encode("utf-8")
    result = document_forensics.analyze_document_file("research.txt", raw_bytes)
    assert result.citations_detected >= 2
    assert len(result.citation_issues) > 0
    assert any("References/Bibliography" in issue for issue in result.citation_issues)

def test_docx_font_shift_detection():
    """Verify DOCX inspection flags copy-paste font family discrepancies."""
    doc = docx.Document()
    p = doc.add_paragraph()
    r1 = p.add_run("This text was authored natively in Calibri. ")
    r1.font.name = "Calibri"
    r2 = p.add_run("This text was pasted from ChatGPT in Segoe UI.")
    r2.font.name = "Segoe UI"
    
    doc.core_properties.author = "Test Author"
    doc.core_properties.revision = 2
    
    bio = io.BytesIO()
    doc.save(bio)
    raw_bytes = bio.getvalue()

    result = document_forensics.analyze_document_file("essay.docx", raw_bytes)
    assert result.file_name == "essay.docx"
    assert result.metadata["author"] == "Test Author"
    assert "Calibri" in result.metadata["detected_font_families"]
    assert "Segoe UI" in result.metadata["detected_font_families"]
    font_anomalies = [a for a in result.formatting_anomalies if a.type == "font"]
    assert len(font_anomalies) > 0
    assert "In-paragraph font family shifts detected" in font_anomalies[0].description

def test_pdf_generation_and_inspection():
    """Verify PDF metadata, page count, and text extraction."""
    writer = pypdf.PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    writer.add_metadata({
        "/Author": "Dr. Forensic Analyst",
        "/Producer": "ReportLab PDF Library",
        "/CreationDate": "D:20260918120000"
    })
    bio = io.BytesIO()
    writer.write(bio)
    raw_bytes = bio.getvalue()

    result = document_forensics.analyze_document_file("report.pdf", raw_bytes)
    assert result.file_name == "report.pdf"
    assert result.pages == 1
    assert result.metadata["author"] == "Dr. Forensic Analyst"
    # Producer "ReportLab" triggers automated compiler signature anomaly
    meta_anomalies = [a for a in result.formatting_anomalies if a.type == "metadata"]
    assert any("compiler signature" in a.description.lower() for a in meta_anomalies)

def test_documents_api_upload_endpoint():
    """Verify POST /api/documents/analyze end-to-end integration."""
    content = b"The dominant sequence transduction models are based on complex recurrent or convolutional neural networks."
    files = {"file": ("test_doc.txt", content, "text/plain")}
    response = client.post("/api/documents/analyze", files=files)
    assert response.status_code == 200
    data = response.json()
    assert "classification" in data
    assert "probability" in data
    assert "analysis_id" in data
    assert "file_name" in data
    assert data["file_name"] == "test_doc.txt"
    assert "text_analysis" in data
    assert data["text_analysis"]["classification"] in ["AI_GENERATED", "AI_ASSISTED", "AUTHENTIC", "UNCERTAIN"]
